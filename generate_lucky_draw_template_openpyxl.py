from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

OUTPUT_FILE = Path("src/frontend/assets/Asta_Studio_Participant_Template.xlsx")

EMERALD = "0F766E"
EMERALD_DARK = "115E59"
SLATE_900 = "0F172A"
SLATE_800 = "1E293B"
SLATE_700 = "334155"
SLATE_600 = "475569"
SLATE_200 = "E2E8F0"
SLATE_100 = "F1F5F9"
SLATE_50 = "F8FAFC"
TEAL_100 = "CCFBF1"
WHITE = "FFFFFF"
YELLOW_200 = "FEF08A"
YELLOW_800 = "854D0E"
GREEN_50 = "ECFDF5"
GREEN_700 = "047857"
RED_50 = "FEF2F2"
RED_700 = "B91C1C"

HEADERS = [
    "ID_Ticket",
    "Full_Name",
    "Department_Company",
    "Phone_Number",
    "Category_Eligibility",
]

CATEGORY_OPTIONS = [
    "All",
    "Grand Prizes",
    "VIP Rounds",
    "Regular Draw",
    "Consolation",
]

SAMPLE_ROWS = [
    ["EMP-001", "Sokha Chan", "Human Resources", "012345678", "All"],
    ["VIP-888", "Dr. Dara Vann", "Executive Office", "+855 12 888 999", "VIP Rounds"],
    ["00123", "Emily Johnson", "Global Partnerships", "+1 415 555 0198", "Grand Prizes"],
    ["SG-042", "Lim Wei Ming", "Singapore Regional Office", "+65 9123 4567", "Regular Draw"],
    ["JP-007", "Haruka Tanaka", "International Sales", "+81 90-1234-5678", "Grand Prizes"],
    ["KH-BR01-015", "Sreyneang Hem", "Phnom Penh Branch", "096 555 0123", "Regular Draw"],
    ["TABLE-A12", "Omar Al-Farsi", "Guest – Table A12", "+971 50 123 4567", "Consolation"],
    ["STAFF-2048", "María García-López", "Marketing & Communications", "+34 612 345 678", "All"],
    ["VIP-CHAIR-01", "Prof. Nguyễn Minh Anh", "Board of Directors", "+84 912 345 678", "VIP Rounds"],
    ["EXPO-0009", "Jean-Luc Moreau", "Expo Partner Pavilion", "+33 6 12 34 56 78", "Consolation"],
]

thin_gray = Side(style="thin", color=SLATE_200)
thin_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

def solid(color):
    return PatternFill(fill_type="solid", fgColor=color)

def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

def style_header(ws, row):
    for col_idx, value in enumerate(HEADERS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.fill = solid(EMERALD)
        cell.font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(
            horizontal="center" if col_idx in (1, 4, 5) else "left",
            vertical="center"
        )
        cell.border = thin_border
    ws.row_dimensions[row].height = 28

def style_data_rows(ws, start_row, end_row):
    for row in range(start_row, end_row + 1):
        ws.row_dimensions[row].height = 22
        row_fill = WHITE if row % 2 == 0 else SLATE_50
        for col_idx in range(1, 6):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = solid(row_fill)
            cell.font = Font(name="Segoe UI", size=10, color=SLATE_900)
            cell.alignment = Alignment(
                horizontal="center" if col_idx in (1, 4, 5) else "left",
                vertical="center"
            )
            cell.border = thin_border

def configure_participant_sheet(ws, start_row, end_row):
    set_widths(ws, {"A": 16, "B": 32, "C": 28, "D": 20, "E": 22})
    for row in range(start_row, end_row + 1):
        ws.cell(row=row, column=1).number_format = "@"
        ws.cell(row=row, column=4).number_format = "@"

    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(CATEGORY_OPTIONS) + '"',
        allow_blank=True
    )
    dv.error = "Select one of the approved eligibility categories."
    dv.errorTitle = "Invalid Category"
    dv.prompt = "Choose the participant's eligible prize tier."
    dv.promptTitle = "Category Eligibility"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"E{start_row}:E{end_row}")

def build_workbook():
    wb = Workbook()
    guide = wb.active
    guide.title = "📖 Instructions & Guide"
    master = wb.create_sheet("👥 Master_Participants_Pool")
    sample = wb.create_sheet("💡 Sample_Data_Examples")

    # TAB 1
    set_widths(guide, {"A": 4, "B": 21, "C": 18, "D": 25, "E": 37, "F": 21, "G": 4})
    guide.sheet_view.showGridLines = False

    guide.merge_cells("A1:G3")
    guide["A1"] = "Asta Studio\nMaster Participant Data Template"
    guide["A1"].fill = solid(SLATE_900)
    guide["A1"].font = Font(name="Segoe UI", size=22, bold=True, color=WHITE)
    guide["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(1, 4):
        guide.row_dimensions[r].height = 28

    guide.merge_cells("B5:F5")
    guide["B5"] = "WELCOME & IMPORT READINESS"
    guide["B5"].fill = solid(TEAL_100)
    guide["B5"].font = Font(name="Segoe UI", size=11, bold=True, color=EMERALD_DARK)
    guide["B5"].alignment = Alignment(horizontal="center", vertical="center")
    guide.row_dimensions[5].height = 24

    guide.merge_cells("B6:F7")
    guide["B6"] = (
        "Use this enterprise-ready workbook to prepare participant data for live gala, "
        "annual dinner, expo, and corporate raffle events. Enter live records only in "
        "the “👥 Master_Participants_Pool” sheet."
    )
    guide["B6"].fill = solid(WHITE)
    guide["B6"].font = Font(name="Segoe UI", size=10, color=SLATE_700)
    guide["B6"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    guide.merge_cells("B9:F9")
    guide["B9"] = "COLUMN DICTIONARY"
    guide["B9"].fill = solid(EMERALD)
    guide["B9"].font = Font(name="Segoe UI", size=12, bold=True, color=WHITE)
    guide["B9"].alignment = Alignment(horizontal="left", vertical="center")
    guide.row_dimensions[9].height = 26

    dictionary_headers = ["Column Key", "Requirement", "Accepted Content", "Description", "Example"]
    for col_idx, value in enumerate(dictionary_headers, start=2):
        cell = guide.cell(row=10, column=col_idx, value=value)
        cell.fill = solid(SLATE_800)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    guide.row_dimensions[10].height = 26

    dictionary_rows = [
        ["ID_Ticket", "Optional*", "Alphanumeric text",
         "Unique employee code, ticket number, badge ID, or draw identifier. Store as text so leading zeros are preserved.",
         "EMP-001 / VIP-888 / 00123"],
        ["Full_Name", "REQUIRED", "Display name",
         "Participant’s complete name exactly as it should appear on the live draw screen.",
         "Sokha Chan"],
        ["Department_Company", "Optional", "Free text",
         "Department, company, branch, business unit, guest group, or table number.",
         "Finance / Branch 01 / Table A12"],
        ["Phone_Number", "Optional", "Text only",
         "Contact number stored strictly as text to preserve leading zeros and international formatting.",
         "012345678 / +855 12 345 678"],
        ["Category_Eligibility", "Optional", "Approved dropdown value",
         "Prize-tier eligibility. Blank values are treated as “All” by the application.",
         "All / VIP Rounds"],
    ]
    for row_idx, values in enumerate(dictionary_rows, start=11):
        for col_idx, value in enumerate(values, start=2):
            cell = guide.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = solid(WHITE if row_idx % 2 else SLATE_50)
            cell.font = Font(name="Segoe UI", size=10, color=SLATE_900)
            cell.alignment = Alignment(
                horizontal="center" if col_idx == 3 else "left",
                vertical="center",
                wrap_text=True
            )
            cell.border = thin_border
        guide.row_dimensions[row_idx].height = 48

    guide["C12"].fill = solid(GREEN_50)
    guide["C12"].font = Font(name="Segoe UI", size=10, bold=True, color=GREEN_700)
    guide["C11"].fill = solid(SLATE_100)
    guide["C11"].font = Font(name="Segoe UI", size=10, bold=True, color=SLATE_600)

    guide.merge_cells("B17:F17")
    guide["B17"] = "BEST PRACTICES & PRO-TIPS"
    guide["B17"].fill = solid(YELLOW_200)
    guide["B17"].font = Font(name="Segoe UI", size=12, bold=True, color=YELLOW_800)
    guide["B17"].alignment = Alignment(horizontal="left", vertical="center")
    guide.row_dimensions[17].height = 26

    tips = [
        "Full_Name is the only mandatory field when the application is drawing by Name.",
        "ID_Ticket becomes mandatory when the application is configured to draw by ID or ticket mode.",
        "Do not rename, translate, reorder, merge, or delete the five column headers in Row 1 of the master sheet.",
        "Keep ID_Ticket and Phone_Number as text. The template is preformatted to preserve values such as 00123 and 012345678.",
        "Use the Category_Eligibility dropdown to prevent spelling differences and invalid prize-tier values.",
        "Before import, remove duplicate participants and confirm that blank rows do not contain hidden spaces.",
    ]
    for row_idx, tip in enumerate(tips, start=18):
        guide.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=6)
        num_cell = guide.cell(row=row_idx, column=2, value=row_idx - 17)
        num_cell.fill = solid(EMERALD)
        num_cell.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border = thin_border

        tip_cell = guide.cell(row=row_idx, column=3, value=tip)
        tip_cell.fill = solid(WHITE if row_idx % 2 == 0 else SLATE_50)
        tip_cell.font = Font(name="Segoe UI", size=10, color=SLATE_800)
        tip_cell.alignment = Alignment(vertical="center", wrap_text=True)
        tip_cell.border = thin_border
        guide.row_dimensions[row_idx].height = 34

    guide.merge_cells("B25:F26")
    guide["B25"] = (
        "IMPORT SAFETY CHECK\n"
        "Save the completed workbook as .xlsx or export the master sheet as UTF-8 CSV. "
        "Never paste formulas into participant fields—use final text values only."
    )
    guide["B25"].fill = solid(RED_50)
    guide["B25"].font = Font(name="Segoe UI", size=10, bold=True, color=RED_700)
    guide["B25"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    red_side = Side(style="thin", color=RED_700)
    guide["B25"].border = Border(left=red_side, right=red_side, top=red_side, bottom=red_side)
    guide.freeze_panes = "A4"

    # TAB 2
    master.sheet_view.showGridLines = False
    style_header(master, 1)
    style_data_rows(master, 2, 50)
    configure_participant_sheet(master, 2, 5000)
    master.freeze_panes = "A2"
    master.auto_filter.ref = "A1:E50"

    # TAB 3
    sample.sheet_view.showGridLines = False
    sample.merge_cells("A1:E1")
    sample["A1"] = (
        "This sheet is for reference/inspiration only. Please enter your live event data "
        "in the 'Master_Participants_Pool' sheet."
    )
    sample["A1"].fill = solid(YELLOW_200)
    sample["A1"].font = Font(name="Segoe UI", size=10, bold=True, color=YELLOW_800)
    sample["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    yellow_side = Side(style="thin", color=YELLOW_800)
    sample["A1"].border = Border(left=yellow_side, right=yellow_side, top=yellow_side, bottom=yellow_side)
    sample.row_dimensions[1].height = 32

    style_header(sample, 3)
    for row_idx, values in enumerate(SAMPLE_ROWS, start=4):
        for col_idx, value in enumerate(values, start=1):
            sample.cell(row=row_idx, column=col_idx, value=value)
    style_data_rows(sample, 4, 13)
    configure_participant_sheet(sample, 4, 200)
    sample.freeze_panes = "A4"
    sample.auto_filter.ref = "A3:E13"

    wb.active = 0
    wb.save(OUTPUT_FILE)

    check = load_workbook(OUTPUT_FILE, read_only=False, data_only=False)
    expected = [
        "📖 Instructions & Guide",
        "👥 Master_Participants_Pool",
        "💡 Sample_Data_Examples",
    ]
    assert check.sheetnames == expected
    assert [check[expected[1]].cell(1, c).value for c in range(1, 6)] == HEADERS
    assert check[expected[1]]["A2"].number_format == "@"
    assert check[expected[1]]["D2"].number_format == "@"
    assert len(check[expected[1]].data_validations.dataValidation) == 1
    check.close()

if __name__ == "__main__":
    build_workbook()
    print(f"Created: {OUTPUT_FILE.resolve()}")
