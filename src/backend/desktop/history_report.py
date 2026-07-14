from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


RESULT_HEADERS = (
    "Draw Time",
    "Round",
    "Prize Category",
    "Draw Source",
    "Slot",
    "Winner",
    "Winner ID",
    "Participant Name",
)

ACTIVITY_HEADERS = (
    "Time",
    "Event",
    "Event ID",
    "Round",
    "Prize Category",
    "Draw Source",
    "Slot",
    "Previous Winner",
    "Previous Winner ID",
    "New Winner",
    "New Winner ID",
    "Winner Count",
    "Details",
)

NAVY = "142536"
NAVY_LIGHT = "20384D"
CYAN = "22B8D8"
CYAN_LIGHT = "DDF6FB"
GREEN_LIGHT = "E8F7EC"
RED_LIGHT = "FDEBEC"
TEXT = "17202A"
MUTED = "607080"
WHITE = "FFFFFF"
GRID = "D8E1E8"
ROW_ALT = "F5F8FA"


def safe_history_filename(project_name: str, scope: str = "all") -> str:
    stem = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', "_", str(project_name or "Asta Studio"))
    stem = re.sub(r"\s+", " ", stem).strip(" ._") or "Asta Studio"
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem = f"_{stem}"
    stem = stem[:80].rstrip(" .") or "Asta Studio"
    suffix = "all" if scope == "all" else "current-view"
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    return f"{stem} - Draw History - {suffix} - {timestamp}.xlsx"


def _safe_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float, bool, datetime)):
        return value
    return str(value)


def _set_cell_value(cell, value: Any) -> Any:
    normalized = _safe_cell_value(value)
    cell.value = normalized
    if isinstance(normalized, str):
        cell.data_type = "s"
    return normalized


def _parse_timestamp(value: Any) -> Any:
    if not isinstance(value, str) or not value.strip():
        return _safe_cell_value(value)
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
        return parsed
    except ValueError:
        return _safe_cell_value(value)


def _solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _configure_page(ws, orientation: str = "landscape") -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.oddHeader.center.text = "&B Asta Studio - Draw History"
    ws.oddFooter.left.text = "Generated &D &T"
    ws.oddFooter.right.text = "Page &P of &N"


def _style_report_header(ws, title: str, subtitle: str, column_count: int) -> None:
    last_column = max(1, column_count)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = ws.cell(1, 1, title)
    title_cell.fill = _solid(NAVY)
    title_cell.font = Font(name="Aptos Display", size=17, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.data_type = "s"
    subtitle_cell.fill = _solid(NAVY_LIGHT)
    subtitle_cell.font = Font(name="Aptos", size=9, color="DCEAF4")
    subtitle_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 20


def _write_data_sheet(
    wb: Workbook,
    name: str,
    title: str,
    subtitle: str,
    rows: Iterable[Dict[str, Any]],
    headers: tuple[str, ...],
) -> int:
    ws = wb.create_sheet(name)
    rows = list(rows)
    _configure_page(ws)
    _style_report_header(ws, title, subtitle, len(headers))

    thin = Side(style="thin", color=GRID)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(4, column, header)
        cell.fill = _solid(CYAN)
        cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 24

    time_headers = {"Draw Time", "Time"}
    identifier_headers = {
        "Winner",
        "Winner ID",
        "Participant Name",
        "Previous Winner",
        "Previous Winner ID",
        "New Winner",
        "New Winner ID",
        "Event ID",
    }
    numeric_headers = {"Round", "Slot", "Winner Count"}
    for row_index, row in enumerate(rows, start=5):
        row_fill = WHITE if row_index % 2 else ROW_ALT
        event_name = str(row.get("Event") or "")
        if name == "Activity" and event_name.startswith("Redraw"):
            row_fill = RED_LIGHT
        elif name == "Activity" and event_name == "Draw":
            row_fill = GREEN_LIGHT

        for column, header in enumerate(headers, start=1):
            raw_value = row.get(header, "")
            value = _parse_timestamp(raw_value) if header in time_headers else _safe_cell_value(raw_value)
            cell = ws.cell(row_index, column)
            _set_cell_value(cell, value)
            cell.fill = _solid(row_fill)
            cell.font = Font(name="Aptos", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=header in {"Details", "Previous Winner", "New Winner"})
            cell.border = Border(bottom=thin)
            if header in time_headers and isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif header in identifier_headers:
                cell.number_format = "@"
            elif header in numeric_headers and value != "":
                cell.number_format = "0"
        ws.row_dimensions[row_index].height = 21

    widths = {
        "Draw Time": 21,
        "Time": 21,
        "Round": 9,
        "Prize Category": 21,
        "Draw Source": 19,
        "Slot": 8,
        "Winner": 25,
        "Winner ID": 19,
        "Participant Name": 25,
        "Event": 17,
        "Event ID": 39,
        "Previous Winner": 28,
        "Previous Winner ID": 20,
        "New Winner": 28,
        "New Winner ID": 20,
        "Winner Count": 13,
        "Details": 40,
    }
    for column, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(4, column).column_letter].width = widths.get(header, 18)

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ws.cell(max(4, len(rows) + 4), len(headers)).coordinate}"
    ws.print_title_rows = "4:4"
    ws.print_area = f"A1:{ws.cell(max(4, len(rows) + 4), len(headers)).coordinate}"

    if rows:
        table = Table(displayName=f"History{name}Table", ref=f"A4:{ws.cell(len(rows) + 4, len(headers)).coordinate}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(headers))
        empty_cell = ws.cell(5, 1)
        _set_cell_value(empty_cell, "No data")
        empty_cell.font = Font(name="Aptos", size=10, italic=True, color=MUTED)
        empty_cell.alignment = Alignment(horizontal="center")

    return len(rows)


def _write_summary_sheet(wb: Workbook, report: Dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "Summary"
    _configure_page(ws, "portrait")
    _style_report_header(ws, "Draw History Report", str(report.get("projectName") or "Untitled Project"), 4)

    summary = report.get("summary") if isinstance(report.get("summary"), dict) else {}
    metric_labels = ("Rounds", "Winners", "Redraws", "Latest Draw")
    metric_values = (
        int(summary.get("rounds") or 0),
        int(summary.get("winners") or 0),
        int(summary.get("redraws") or 0),
        str(summary.get("latest") or "--"),
    )
    for column, (label, value) in enumerate(zip(metric_labels, metric_values), start=1):
        label_cell = ws.cell(4, column, label)
        label_cell.fill = _solid(CYAN_LIGHT)
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=MUTED)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell = ws.cell(5, column)
        _set_cell_value(value_cell, value)
        value_cell.fill = _solid(WHITE)
        value_cell.font = Font(name="Aptos Display", size=15, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 32

    filters = report.get("filters") if isinstance(report.get("filters"), dict) else {}
    metadata = (
        ("Project", report.get("projectName") or "Untitled Project"),
        ("Scope", "All History" if report.get("scope") == "all" else "Current View"),
        ("View", str(report.get("view") or "results").title()),
        ("Search", filters.get("search") or "All"),
        ("Category", filters.get("categoryLabel") or "All Categories"),
        ("Draw Source", filters.get("sourceLabel") or "All Sources"),
        ("Sort", filters.get("sortLabel") or "Newest First"),
        ("Generated", _parse_timestamp(report.get("generatedAt"))),
    )
    start_row = 8
    for offset, (label, value) in enumerate(metadata):
        row = start_row + offset
        ws.cell(row, 1, label).font = Font(name="Aptos", size=10, bold=True, color=MUTED)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        value_cell = ws.cell(row, 2)
        _set_cell_value(value_cell, value)
        if label == "Generated" and isinstance(value, datetime):
            value_cell.value = value
            value_cell.number_format = "yyyy-mm-dd hh:mm:ss"
        value_cell.font = Font(name="Aptos", size=10, color=TEXT)
        value_cell.alignment = Alignment(horizontal="left", wrap_text=True)

    ws.column_dimensions["A"].width = 18
    for column in ("B", "C", "D"):
        ws.column_dimensions[column].width = 20
    ws.print_area = f"A1:D{start_row + len(metadata) - 1}"


def build_history_workbook(report: Dict[str, Any]) -> tuple[bytes, Dict[str, Any]]:
    if not isinstance(report, dict):
        raise ValueError("History report must be an object")

    results = report.get("results") if isinstance(report.get("results"), list) else []
    activity = report.get("activity") if isinstance(report.get("activity"), list) else []
    scope = "all" if report.get("scope") == "all" else "view"
    view = "activity" if report.get("view") == "activity" else "results"
    if not results and not activity:
        raise ValueError("No history data is available to export")

    wb = Workbook()
    wb.properties.title = "Asta Studio - Draw History"
    wb.properties.subject = "Draw results and activity audit"
    wb.properties.creator = "Asta Studio"
    _write_summary_sheet(wb, report)

    generated_label = str(report.get("generatedAt") or "")
    project_label = str(report.get("projectName") or "Untitled Project")
    row_counts: Dict[str, int] = {}
    if scope == "all" or view == "results":
        row_counts["Results"] = _write_data_sheet(
            wb, "Results", "Draw Results", f"{project_label} | Generated {generated_label}", results, RESULT_HEADERS
        )
    if scope == "all" or view == "activity":
        row_counts["Activity"] = _write_data_sheet(
            wb, "Activity", "Draw Activity", f"{project_label} | Generated {generated_label}", activity, ACTIVITY_HEADERS
        )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), {"sheets": list(wb.sheetnames), "rowCounts": row_counts}


def write_history_workbook(path: Path, report: Dict[str, Any]) -> Dict[str, Any]:
    output, metadata = build_history_workbook(report)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: str | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, dir=path.parent, suffix=".xlsx.tmp") as temporary:
            temporary.write(output)
            temporary.flush()
            os.fsync(temporary.fileno())
            temporary_path = temporary.name
        os.replace(temporary_path, path)
    finally:
        if temporary_path:
            Path(temporary_path).unlink(missing_ok=True)
    return {"ok": True, "path": str(path), **metadata}
